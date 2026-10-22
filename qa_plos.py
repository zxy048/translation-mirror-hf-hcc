# -*- coding: utf-8 -*-
# HISTORICAL (v1) CHECKER — SUPERSEDED. Kept as the worked example described in
# Code/README.md and archive/README.md. It targets Manuscript_PLOSONE_v1.md and
# Manuscript_PLOSONE_v1.docx by absolute path and assumes the v1 section layout
# and reference count. It is NOT the current QA entry point and will not run in
# a clean clone without those v1 files; run verify_all.py for the revision.
"""
QA for PLOS ONE manuscript (md + docx):
  md-level:
    - abstract <= 300 words, unstructured, no citations
    - every [N] citation appears ascending within groups, numbered by first appearance
    - all 26 refs cited; no leftover Harvard (et al., YYYY)
    - figure captions Fig 1-5 present inline; supp fig captions section present
    - required PLOS sections present
  docx-level:
    - Letter page size, 1-in margins
    - double spacing (2.0) on body
    - continuous line numbers (w:lnNumType)
    - footer page number field
"""
import re
import sys

MD = r"D:\R_projects\revision_analysis\Manuscript_PLOSONE_v1.md"
DOCX = r"D:\R_projects\revision_analysis\Manuscript_PLOSONE_v1.docx"

fail = 0
def ok(cond, msg):
    global fail
    print(("  PASS  " if cond else "  FAIL  ") + msg)
    if not cond:
        fail += 1

with open(MD, encoding="utf-8") as f:
    md = f.read()

# ---- md-level ----
body = md.split("## References")[0]  # text before references (includes abstract)

abs_start = md.index("## Abstract")
abs_end = md.index("## Introduction")
abstract = md[abs_start:abs_end]
abs_words = [w for w in re.sub(r"[#*\n]", " ", abstract).split() if w.strip()]
print(f"Abstract word count: {len(abs_words)} (limit 300)")
ok(len(abs_words) <= 300, f"abstract <= 300 words ({len(abs_words)})")
ok("[" not in abstract.replace("## Abstract", ""), "abstract contains no citations")

# citation groups -> first-appearance order check
# (only integer tokens 1-26; excludes numeric CIs like [0.67, 4.21])
cites = []
for m in re.finditer(r"\[([0-9]+(?:,[0-9]+)*)\]", body):
    nums = [int(t) for t in m.group(1).split(",")]
    if all(1 <= n <= 26 for n in nums):
        cites.append(nums)
first_seen = {}
order = []
for group in cites:
    for n in group:
        if n not in first_seen:
            first_seen[n] = len(order)
            order.append(n)
ok(order == sorted(order) and len(set(order)) == 26,
   f"citations numbered by first appearance (first-order = ref numbering); {len(order)} refs cited")
ok(all(g == sorted(g) for g in cites), "no descending citation groups (e.g. [8,5,7])")

# all 26 refs cited
ref_block = md[md.index("## References"):md.index("## Supporting Information Captions")]
refs = re.findall(r"^(\d+)\.\s", ref_block, flags=re.M)
ok(len(refs) == 26, f"reference list has 26 entries (got {len(refs)})")
ok(set(int(r) for r in refs) == set(first_seen.keys()),
   "reference numbering matches first-appearance order")

# no leftover Harvard citations
leftover = re.findall(r"et al\.,\s+\d{4}|\([A-Z][a-z]+ et al\.|\([A-Z][a-z]+ and [A-Z][a-z]+, \d{4}", body)
ok(not leftover, f"no leftover Harvard citations ({len(leftover)})")

# figure captions inline
for n in range(1, 6):
    m = re.search(rf"\) Fig {n}[ab]?(?:,|\))|\bFig {n}[ab]?\b", body) if False else re.search(
        r"\bFig %d\b" % n, body)
    ok(bool(m), f"Fig {n} cited in text")
    cap = re.search(rf"\*\*Fig {n}\.\*\*", body)
    ok(bool(cap), f"Fig {n} caption present inline")

# required sections
for sec in ["## Abstract", "## Introduction", "## Materials and Methods",
            "## Results", "## Discussion", "## References",
            "## Supporting Information Captions"]:
    ok(sec in md, f"section present: {sec}")

# ---- docx-level ----
from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn
doc = Document(DOCX)

# page size + margins
sec = doc.sections[0]
W = round(sec.page_width.cm, 1); H = round(sec.page_height.cm, 1)
ok((W, H) == (21.6, 27.9), f"page size Letter ({W:.1f} x {H:.1f} cm)")
ok(round(sec.left_margin.cm, 2) == 2.54 and round(sec.right_margin.cm, 2) == 2.54,
   "left/right margins 1 in (2.54 cm)")

# double spacing on body paragraphs
double_spaced = sum(1 for p in doc.paragraphs
                    if p.text.strip() and p.paragraph_format.line_spacing == 2.0)
total_paras = sum(1 for p in doc.paragraphs if p.text.strip())
ok(double_spaced / max(total_paras, 1) > 0.9,
   f"double spacing on body ({double_spaced}/{total_paras} non-empty paras)")

# line numbers present
ln = sec._sectPr.find(qn('w:lnNumType'))
ok(ln is not None and ln.get(qn('w:restart')) == 'continuous',
   "continuous line numbers present")

# footer page number field
footer_xml = sec.footer.paragraphs[0]._p.xml if sec.footer.paragraphs else ""
ok("PAGE" in footer_xml, "footer page-number field present")

# abstract on new page
abs_para = next((p for p in doc.paragraphs if p.text.strip() == "Abstract"), None)
ok(abs_para is not None and bool(abs_para.paragraph_format.page_break_before),
   "Abstract starts on a new page")

print("\n" + ("ALL CHECKS PASSED" if fail == 0 else f"{fail} CHECK(S) FAILED"))
sys.exit(1 if fail else 0)
