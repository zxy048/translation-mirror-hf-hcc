# -*- coding: utf-8 -*-
"""Apply the 2026-10-22 language and layout patch to the response letter.

The response letter is the package artifact of record; this script edits that
artifact directly and records the before/after state. It is idempotent: a
second run reports "already applied" and changes nothing.

Three text replacements (each must occur exactly once):

  L1  the sentence that gives the two DCM/ICM mean effect sizes now states
      them numerically instead of pointing at the comment that reported them.
  L2  "cited to" -> "cross-referenced to", and "each of them" -> "each", so
      the sentence matches the table it names.
  L3  removes the duplicated "with" before the small-sample correction.

Two layout items:

  L4  table column widths are set explicitly (fixed layout, 30-twip left and
      right cell margins, grids summing to 15.1 cm on a 15.24 cm text width).
      The widths were derived from the rendered PDF so that no table token is
      broken inside a word. Cell text is not modified.
  L5  a standalone italic "P" is set upright. The rendered letter has 26
      upright and 5 italic instances of the statistic symbol; the majority
      convention is upright, so all five are unified to it.

Guards:

  * every L1-L3 old string must occur exactly once before the patch; the new
    text must occur at least once afterwards (L1's numeric form already
    appears in an earlier paragraph, and the patch makes the two consistent);
    replacing a run must not cross runs;
  * the four reviewer-opinion blocks are hashed before and after and must be
    byte-identical (their text is never stored in this script);
  * all paragraph and table-cell text outside the three target paragraphs must
    be unchanged;
  * table count, column counts and cell text are asserted unchanged.

This script writes only the DOCX. If `assemble_plos_v2.py` is ever rerun, it
regenerates the letter from Markdown and this patch must be applied again
before the package hash is regenerated.
"""
import hashlib
import os
import shutil
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DOCX = os.path.join(ROOT, "PLOS_ONE_Submission_v2",
                    "02_Response_to_Reviewers.docx")

# (label, old, new)
REPLACEMENTS = [
    ("L1",
     "although their mean effect sizes differed (Major comment 3)",
     "although their mean effect sizes differed (mean g = −0.291 and −0.431)"),
    ("L2",
     "the correlations are cited to the per-cohort effect-size table (S14) "
     "from which each of them is computed",
     "the correlations are cross-referenced to the per-cohort effect-size "
     "table (S14), from which each is computed"),
    ("L3",
     "with unequal group sizes handled by the weighting, with the "
     "small-sample correction J applied.",
     "with unequal group sizes handled by the weighting and the small-sample "
     "correction J applied."),
]

# Column widths in cm; each table's grid sums to 15.1 cm. Derived from the
# rendered PDF: every token of six characters or more is intact at these
# widths, including accession-like tokens containing digits and underscores.
WIDTHS = {
    0: [2.00, 1.80, 1.35, 1.40, 2.20, 2.60, 1.30, 2.45],
    1: [4.00, 2.775, 2.775, 2.775, 2.775],
    2: [2.30, 2.30, 2.20, 3.70, 2.60, 2.00],
    3: [2.60, 1.20, 1.30, 1.50, 1.50, 1.60, 1.60, 2.10, 1.70],
    4: [3.80, 1.20, 2.80, 3.30, 4.00],
    5: [2.60, 2.20, 1.50, 2.60, 2.60, 3.60],
    6: [1.00, 4.90, 4.90, 4.30],
}
CELL_MARGIN_TWIPS = 30
EXPECTED_QUOTE_PARAGRAPHS = 41

# The paragraph that carried L1 is the only place with this exact context;
# paragraph 144 already stated the same numbers in a different sentence.
L1_CONTEXT = ("each met the mirror criteria for 11 of 13 evaluable sets, "
              "although their mean effect sizes differed "
              "(mean g = −0.291 and −0.431)")


def sha(path):
    with open(path, "rb") as fh:
        return hashlib.sha256(fh.read()).hexdigest()


def count_body(doc, needle):
    return sum(p.text.count(needle) for p in doc.paragraphs)


def replace_once(doc, old, new):
    """Replace one body-paragraph occurrence wholly inside a single run."""
    for p in doc.paragraphs:
        if old not in p.text:
            continue
        for run in p.runs:
            if old in run.text:
                if run.text.count(old) != 1:
                    return -2
                run.text = run.text.replace(old, new, 1)
                return 1
        return -3  # the string spans run boundaries; refuse to touch it
    return 0


def reviewer_quote_hash(doc):
    """SHA256 over the quoted reviewer-opinion paragraphs, in document order.

    The text is hashed, never printed or stored, so this script carries no
    reviewer material. Only paragraphs in "Block Text" style between a
    "Reviewer N" heading and the next Heading 1 are counted.
    """
    h = hashlib.sha256()
    inside = False
    n = 0
    for p in doc.paragraphs:
        if p.style.name == "Heading 1":
            inside = p.text.strip().startswith("Reviewer ")
        elif inside and p.style.name == "Block Text":
            h.update(p.text.encode("utf-8"))
            h.update(b"\n")
            n += 1
    return h.hexdigest(), n


def set_cell_margins(table, twips):
    tblPr = table._tbl.tblPr
    old = tblPr.find(qn("w:tblCellMar"))
    if old is not None:
        tblPr.remove(old)
    mar = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + side)
        e.set(qn("w:w"), "0" if side in ("top", "bottom") else str(twips))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)


def apply_widths(doc):
    if len(doc.tables) != len(WIDTHS):
        raise SystemExit("ABORT: expected %d tables, found %d"
                         % (len(WIDTHS), len(doc.tables)))
    for ti, table in enumerate(doc.tables):
        widths = WIDTHS[ti]
        if len(table.columns) != len(widths):
            raise SystemExit("ABORT: table %d has %d columns, expected %d"
                             % (ti, len(table.columns), len(widths)))
        if abs(sum(widths) - 15.10) > 0.02:
            raise SystemExit("ABORT: table %d widths sum to %.3f cm"
                             % (ti, sum(widths)))

        tblPr = table._tbl.tblPr
        for tag in ("w:tblW", "w:tblLayout"):
            old = tblPr.find(qn(tag))
            if old is not None:
                tblPr.remove(old)
        tblW = OxmlElement("w:tblW")
        tblW.set(qn("w:w"), "5000")
        tblW.set(qn("w:type"), "pct")
        tblPr.append(tblW)
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        set_cell_margins(table, CELL_MARGIN_TWIPS)

        grid = table._tbl.tblGrid
        if grid is None:
            grid = OxmlElement("w:tblGrid")
            table._tbl.insert(1, grid)
        for g in list(grid):
            grid.remove(g)
        for w in widths:
            gc = OxmlElement("w:gridCol")
            gc.set(qn("w:w"), str(int(round(w * 567))))
            grid.append(gc)
        for row in table.rows:
            for ci, cell in enumerate(row.cells):
                cell.width = Cm(widths[ci])


def grid_signature(doc):
    sig = []
    for table in doc.tables:
        grid = table._tbl.tblGrid
        sig.append([int(g.get(qn("w:w"))) for g in grid])
    return sig


def target_grids():
    return [[int(round(w * 567)) for w in WIDTHS[i]] for i in sorted(WIDTHS)]


def italic_p_runs(doc):
    out = []
    for p in doc.paragraphs:
        for r in p.runs:
            if r.italic and r.text.strip() == "P":
                out.append(r)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if r.italic and r.text.strip() == "P":
                            out.append(r)
    return out


def main():
    if not os.path.exists(DOCX):
        sys.exit("ABORT: document not found: %s" % DOCX)

    before_hash = sha(DOCX)
    doc = Document(DOCX)

    para_before = [p.text for p in doc.paragraphs]
    cells_before = [[c.text for c in row.cells]
                    for t in doc.tables for row in t.rows]
    quote_before, quote_n = reviewer_quote_hash(doc)
    if quote_n != EXPECTED_QUOTE_PARAGRAPHS:
        sys.exit("ABORT: found %d reviewer-quote paragraphs, expected %d"
                 % (quote_n, EXPECTED_QUOTE_PARAGRAPHS))
    grids_before = grid_signature(doc)
    it_before = len(italic_p_runs(doc))

    # Already patched? All three old strings absent. L1's new numeric form
    # already existed in another paragraph, so the check is per old string.
    text_applied = all(count_body(doc, old) == 0 for _, old, _ in REPLACEMENTS)
    if text_applied:
        print("text already applied: all three old strings absent")
        for label, _, new in REPLACEMENTS:
            if count_body(doc, new) < 1:
                sys.exit("ABORT: %s new text absent although old text is gone"
                         % label)
        if count_body(doc, L1_CONTEXT) != 1:
            sys.exit("ABORT: L1 target paragraph not found in patched form")
    else:
        for label, old, new in REPLACEMENTS:
            n_old = count_body(doc, old)
            if n_old != 1:
                sys.exit("ABORT: %s old occurs %d times (expected 1)"
                         % (label, n_old))
        for label, old, new in REPLACEMENTS:
            rc = replace_once(doc, old, new)
            if rc != 1:
                sys.exit("ABORT: %s replacement failed (code %d)" % (label, rc))
            print("%s: replaced one occurrence" % label)
        for label, old, new in REPLACEMENTS:
            if count_body(doc, old) != 0 or count_body(doc, new) < 1:
                sys.exit("ABORT: %s post-check failed" % label)
        if count_body(doc, L1_CONTEXT) != 1:
            sys.exit("ABORT: L1 target paragraph did not acquire the numeric form")

    # Every other paragraph must be text-identical.
    para_after = [p.text for p in doc.paragraphs]
    if len(para_after) != len(para_before):
        sys.exit("ABORT: paragraph count changed")
    expected = list(para_before)
    for _, old, new in REPLACEMENTS:
        for i, text in enumerate(expected):
            if old in text:
                expected[i] = text.replace(old, new, 1)
                break
    if para_after != expected:
        sys.exit("ABORT: unexpected paragraph text change")

    # Reviewer quotes and table cell text must not move.
    quote_after, n_after = reviewer_quote_hash(doc)
    if n_after != quote_n or quote_after != quote_before:
        sys.exit("ABORT: reviewer-quote blocks changed")
    if [[c.text for c in row.cells] for t in doc.tables for row in t.rows] != cells_before:
        sys.exit("ABORT: table cell text changed")

    if text_applied and grids_before == target_grids() and it_before == 0:
        print("already applied: text, table widths and italic P are all at target")
        print("no change made")
        return

    # L4 layout and L5 italic unification.
    apply_widths(doc)
    grids_after = grid_signature(doc)
    if grids_after != target_grids():
        sys.exit("ABORT: table grids do not match the target widths")
    italics = italic_p_runs(doc)
    if len(italics) not in (0, 5):
        sys.exit("ABORT: found %d italic standalone P runs, expected 5 or 0"
                 % len(italics))
    for run in italics:
        run.italic = False
    it_after = len(italic_p_runs(doc))

    tmp = DOCX + ".r33tmp"
    doc.save(tmp)
    os.replace(tmp, DOCX)

    after_hash = sha(DOCX)
    if after_hash == before_hash:
        print("no content change: document hash unchanged")
    print("document sha256 %s -> %s" % (before_hash, after_hash))
    print("reviewer-quote paragraphs: %d, sha256 %s (unchanged)"
          % (quote_n, quote_before))
    print("table grids before: %s" % grids_before)
    print("table grids after:  %s" % grids_after)
    print("italic standalone P runs: %d -> %d" % (it_before, it_after))
    print("cell text unchanged; all paragraph text outside L1-L3 unchanged")
    print("OK: L1-L5 applied")


if __name__ == "__main__":
    main()
